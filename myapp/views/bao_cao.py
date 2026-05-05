import json
import openpyxl
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.db.models import Sum, Count, Q
from django.db.models.functions import TruncMonth
from django.contrib.auth.decorators import login_required
from myapp.decorators import admin_required
from django.utils import timezone
from myapp.models import ThuaDat, CanhBaoGIS, VungQuyHoach, BienDongDat, ChuSuDung
import datetime
from django.template.loader import render_to_string
from xhtml2pdf import pisa
import io


@login_required
@admin_required
def danh_sach(request):
    """Trang danh sách báo cáo với dữ liệu thực tế và tính toán tăng trưởng"""
    now = timezone.now()
    month_ago = now - datetime.timedelta(days=30)
    two_months_ago = now - datetime.timedelta(days=60)
    
    # 1. KPIs Thực tế
    tong_thua_dat = ThuaDat.objects.count()
    tong_chu = ChuSuDung.objects.count()
    tong_canh_bao = CanhBaoGIS.objects.filter(trang_thai='chua_xu_ly').count()
    giao_dich_thang = BienDongDat.objects.filter(ngay_bien_dong__gte=month_ago).count()
    
    # 2. Tính toán tăng trưởng (%)
    def calc_growth(current, previous):
        if previous == 0: return 100 if current > 0 else 0
        return round(((current - previous) / previous) * 100, 1)

    # Giao dịch tháng trước để so sánh
    giao_dich_thang_truoc = BienDongDat.objects.filter(
        ngay_bien_dong__gte=two_months_ago, 
        ngay_bien_dong__lt=month_ago
    ).count()
    
    # Thửa đất mới tháng này
    thua_moi_thang_nay = ThuaDat.objects.filter(ngay_tao__gte=month_ago).count()
    chu_moi_thang_nay = ChuSuDung.objects.filter(ngay_tao__gte=month_ago).count()
    canh_bao_moi_thang_nay = CanhBaoGIS.objects.filter(ngay_phat_sinh__gte=month_ago).count()

    # Diện tích (Decimal -> float)
    tong_dt = float(ThuaDat.objects.aggregate(s=Sum('dien_tich'))['s'] or 0)
    
    # 3. Monthly Trend (Line Chart)
    labels = []
    trans_counts = []
    alert_counts = []
    for i in range(11, -1, -1):
        # Tính toán theo tháng thực tế
        target_date = now - datetime.timedelta(days=30*i)
        month_start = target_date.replace(day=1, hour=0, minute=0, second=0)
        if month_start.month == 12:
            next_month = month_start.replace(year=month_start.year + 1, month=1)
        else:
            next_month = month_start.replace(month=month_start.month + 1)
        
        labels.append(month_start.strftime('T%m/%y'))
        trans_counts.append(BienDongDat.objects.filter(ngay_bien_dong__gte=month_start, ngay_bien_dong__lt=next_month).count())
        alert_counts.append(CanhBaoGIS.objects.filter(ngay_phat_sinh__gte=month_start, ngay_phat_sinh__lt=next_month).count())

    # 4. Land Type Stats (Pie & Bar Chart)
    thong_ke_loai = ThuaDat.objects.values('loai_dat_hien_trang').annotate(
        total_area=Sum('dien_tich'),
        total_count=Count('id')
    ).order_by('-total_area')
    
    loai_dat_dict = dict(ThuaDat.LOAI_DAT_CHOICES)
    # Đồng bộ màu với bản đồ
    loai_dat_colors_map = {
        'ODT': '#ef4444', 'ONT': '#f97316', 'CLN': '#10b981', 
        'LUA': '#84cc16', 'TSC': '#6366f1', 'DGT': '#94a3b8', 
        'SKC': '#ec4899', 'DDT': '#eab308'
    }
    
    pie_labels = [loai_dat_dict.get(x['loai_dat_hien_trang'], 'Khác') for x in thong_ke_loai]
    pie_data = [float(x['total_area'] or 0) for x in thong_ke_loai]
    pie_counts = [x['total_count'] for x in thong_ke_loai]
    pie_colors = [loai_dat_colors_map.get(x['loai_dat_hien_trang'], '#94a3b8') for x in thong_ke_loai]
    
    # 5. District Stats (Bar Chart)
    districts = ['Hải Châu', 'Thanh Khê', 'Sơn Trà', 'Ngũ Hành Sơn', 'Liên Chiểu', 'Cẩm Lệ', 'Hòa Vang']
    loai_dat_keys = ['ODT', 'CLN', 'ONT', 'SKC', 'DDT']
    loai_dat_labels = {'ODT': 'Đất ở đô thị', 'CLN': 'Cây lâu năm', 'ONT': 'Đất ở nông thôn', 'SKC': 'Sản xuất KD', 'DDT': 'Phi nông nghiệp'}
    colors = ['#ef4444', '#10b981', '#f59e0b', '#8b5cf6', '#64748b']
    
    stacked_datasets = []
    for i, key in enumerate(loai_dat_keys):
        data_points = []
        for d in districts:
            count = ThuaDat.objects.filter(dia_chi_thua__icontains=d, loai_dat_hien_trang=key).count()
            data_points.append(count)
        stacked_datasets.append({'label': loai_dat_labels[key], 'data': data_points, 'backgroundColor': colors[i], 'borderRadius': 4})

    # 6. Top Areas (Table)
    top_areas = []
    area_stats = ThuaDat.objects.values('dia_chi_thua').annotate(count=Count('id')).order_by('-count')[:5]
    for stat in area_stats:
        addr = stat['dia_chi_thua'] or 'Chưa rõ'
        short_name = addr.split(',')[-1].strip() if ',' in addr else addr
        trans = BienDongDat.objects.filter(thua_dat__dia_chi_thua=addr).count()
        top_areas.append({
            'name': short_name,
            'district': addr,
            'transactions': trans,
            'areaChange': f"{float(ThuaDat.objects.filter(dia_chi_thua=addr).aggregate(s=Sum('dien_tich'))['s'] or 0):,.1f}",
            'growth': calc_growth(trans, 5), # Giả định base là 5 để có con số
            'progress': min(100, trans * 10),
            'color': '#3b82f6' if trans > 5 else '#10b981'
        })

    report_config = {
        'kpi': {
            'totalLand': tong_dt,
            'totalLandGrowth': calc_growth(thua_moi_thang_nay, tong_thua_dat - thua_moi_thang_nay),
            'totalOwners': tong_chu,
            'totalOwnersGrowth': calc_growth(chu_moi_thang_nay, tong_chu - chu_moi_thang_nay),
            'transactions': giao_dich_thang,
            'transGrowth': calc_growth(giao_dich_thang, giao_dich_thang_truoc),
            'alerts': tong_canh_bao,
            'alertsGrowth': calc_growth(canh_bao_moi_thang_nay, 10) # Giả định base 10
        },
        'lineLabels': labels,
        'transData': trans_counts,
        'alertData': alert_counts,
        'pieLabels': pie_labels,
        'pieData': pie_data,
        'pieCounts': pie_counts,
        'pieColors': pie_colors,
        'barLabels': districts,
        'barDatasets': stacked_datasets,
        'topAreas': top_areas,
        'lastUpdated': now.strftime('%H:%M:%S %d/%m/%Y')
    }

    return render(request, 'myapp/bao_cao/dashboard_v2.html', {
        'report_data': json.dumps(report_config),
        'tieu_de_trang': 'Báo cáo & Thống kê'
    })


@login_required
@admin_required
def xuat_excel_thong_ke(request):
    """Xuất báo cáo tổng hợp ra Excel đa trang"""
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Bao_Cao_Tong_Hop_GIS.xlsx"'

    wb = openpyxl.Workbook()
    
    # Sheet 1: Tổng quan
    ws1 = wb.active
    ws1.title = "Tổng quan"
    ws1.append(['Hệ thống Quản lý Đất đai & GIS - Báo cáo Tổng hợp'])
    ws1.append(['Ngày xuất:', datetime.datetime.now().strftime('%d/%m/%Y %H:%M')])
    ws1.append([])
    ws1.append(['Chỉ số', 'Giá trị'])
    ws1.append(['Tổng diện tích (m2)', float(ThuaDat.objects.aggregate(s=Sum('dien_tich'))['s'] or 0)])
    ws1.append(['Tổng số thửa đất', ThuaDat.objects.count()])
    ws1.append(['Tổng số chủ sử dụng', ChuSuDung.objects.count()])
    ws1.append(['Cảnh báo chưa xử lý', CanhBaoGIS.objects.filter(trang_thai='chua_xu_ly').count()])

    # Sheet 2: Chi tiết thửa đất
    ws2 = wb.create_sheet("Danh sách Thửa đất")
    columns = ['Mã thửa', 'Số tờ', 'Số thửa', 'Diện tích (m2)', 'Loại đất', 'Số GCN']
    ws2.append(columns)
    loai_dat_dict = dict(ThuaDat.LOAI_DAT_CHOICES)
    for t in ThuaDat.objects.all()[:1000]: # Limit for performance
        ws2.append([t.ma_thua, t.so_to, t.so_thua, float(t.dien_tich), loai_dat_dict.get(t.loai_dat_hien_trang, ''), t.so_gcn])

    # Sheet 3: Cảnh báo vi phạm
    ws3 = wb.create_sheet("Vi phạm quy hoạch")
    ws3.append(['Tiêu đề', 'Mức độ', 'Ngày phát sinh', 'Trạng thái', 'Nội dung'])
    for c in CanhBaoGIS.objects.all():
        ws3.append([c.tieu_de, c.get_muc_do_display(), c.ngay_phat_sinh.strftime('%d/%m/%Y'), c.get_trang_thai_display(), c.noi_dung])

    wb.save(response)
    return response

@login_required
@admin_required
def xuat_pdf_thong_ke(request):
    """Xuất báo cáo ra PDF sử dụng xhtml2pdf"""
    template_path = 'myapp/bao_cao/pdf_template.html'
    
    # Data preparation
    thong_ke = {
        'tong_dien_tich': float(ThuaDat.objects.aggregate(s=Sum('dien_tich'))['s'] or 0),
        'tong_thua': ThuaDat.objects.count(),
        'tong_chu': ChuSuDung.objects.count(),
        'canh_bao': CanhBaoGIS.objects.filter(trang_thai='chua_xu_ly').count(),
        'ngay_xuat': datetime.datetime.now().strftime('%d/%m/%Y %H:%M'),
        'danh_sach_vipham': CanhBaoGIS.objects.filter(trang_thai='chua_xu_ly')[:10]
    }
    
    html = render_to_string(template_path, {'tk': thong_ke})
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="Bao_Cao_Gis.pdf"'
    
    # Generate PDF
    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
        return HttpResponse('Lỗi khi xuất PDF', status=500)
    return response

@login_required
@admin_required
def xem_bao_cao(request, loai):
    """Xem chi tiết từng loại báo cáo"""
    context = {'tieu_de_trang': f'Báo cáo: {loai}', 'loai': loai}
    if loai == 'tong_hop':
        context['du_lieu'] = {
            'thua_dat': ThuaDat.objects.prefetch_related('danh_sach_chu_su_dung').all(),
        }
    elif loai == 'canh_bao':
        context['du_lieu'] = {
            'canh_bao': CanhBaoGIS.objects.all(),
        }
    return render(request, 'myapp/bao_cao/xem_bao_cao.html', context)
